from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
import subprocess
import os
from pdf2image import convert_from_path

def generate_presentation(slide_contents, pptx_path, config=None):
    print("Creating enhanced slides...")
    prs = Presentation()

    if hasattr(slide_contents, 'slides'):
        slides = slide_contents.slides
        theme_colors = slide_contents.theme_colors if slide_contents.theme_colors else {}

    elif isinstance(slide_contents, list) and all(hasattr(item, 'title') for item in slide_contents):
        slides = slide_contents
        theme_colors = {}

    elif isinstance(slide_contents, list) and len(slide_contents) > 0 and 'slides' in dir(slide_contents[0]):
        slides = slide_contents[0].slides
        theme_colors = slide_contents[0].theme_colors if slide_contents[0].theme_colors else {}

    else:
        slides = slide_contents
        theme_colors = {}
        if hasattr(slide_contents, 'theme_colors'):
            theme_colors = slide_contents.theme_colors

    if not theme_colors:
        theme_colors = {
            "primary": "#1F497D",
            "secondary": "#4F81BD",
            "accent": "#C0504D",
            "background": "#FFFFFF",
            "text": "#000000"
        }

    for key in theme_colors:
        if isinstance(theme_colors[key], str) and theme_colors[key].startswith('#'):
            theme_colors[key] = theme_colors[key][1:]

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = title_slide.shapes.title
    subtitle = title_slide.placeholders[1]

    presentation_title = "Presentation"
    if slides and hasattr(slides[0], 'title'):
        first_title = slides[0].title
        if "Introduction" in first_title and "to" in first_title:
            presentation_title = first_title.split("to")[1].strip()
        else:
            presentation_title = first_title

    title.text = presentation_title
    subtitle.text = "A Comprehensive Guide"
    content_about = title.text

    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme_colors["primary"])
    subtitle.text_frame.paragraphs[0].font.size = Pt(28)
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme_colors["secondary"])

    for i, slide_item in enumerate(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title_shape = slide.shapes.title
        content_placeholder = slide.placeholders[1]

        slide_title = f"Slide {i+1}"
        if hasattr(slide_item, 'title'):
            slide_title = slide_item.title

        title_shape.text = slide_title
        title_shape.text_frame.paragraphs[0].font.size = Pt(36)
        title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme_colors["primary"])

        content_frame = content_placeholder.text_frame
        content_frame.clear()

        slide_content = "Content"
        if hasattr(slide_item, 'content'):
            slide_content = slide_item.content

        p = content_frame.add_paragraph()
        p.text = slide_content
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor.from_string(theme_colors["text"])

        key_points = []
        if hasattr(slide_item, 'key_points'):
            key_points = slide_item.key_points

        if key_points:
            content_frame.add_paragraph().text = ""
            for point in key_points:
                bullet_p = content_frame.add_paragraph()
                bullet_p.text = point
                bullet_p.font.size = Pt(20)
                bullet_p.level = 1
                bullet_p.font.color.rgb = RGBColor.from_string(theme_colors["secondary"])

    final_slide = prs.slides.add_slide(prs.slide_layouts[2])
    final_title = final_slide.shapes.title
    final_content = final_slide.placeholders[1]

    final_title.text = "Thank You!"
    final_title.text_frame.paragraphs[0].font.size = Pt(40)
    final_title.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string(theme_colors["primary"])

    final_p = final_content.text_frame.add_paragraph()
    final_p.text = "Please Like, Share And Subscribe?"
    final_p.font.size = Pt(32)
    final_p.font.color.rgb = RGBColor.from_string(theme_colors["accent"])

    prs.save(pptx_path)
    print(f"Enhanced slides created and saved to {pptx_path}")
    return content_about


def slides_to_images(ppt_path, output_folder, slides_data=None, theme_colors=None):
    import platform, shutil
    soffice = shutil.which("soffice")
    if soffice is None:
        if platform.system() == "Windows":
            candidates = [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            ]
            soffice = next((c for c in candidates if os.path.exists(c)), None)
        elif platform.system() == "Darwin":
            soffice = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

    if soffice is not None:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", ppt_path, "--outdir", output_folder], check=True)
        pdf_path = os.path.join(output_folder, os.path.splitext(os.path.basename(ppt_path))[0] + ".pdf")
        return [img.save(os.path.join(output_folder, f"slide_{i}.png"), 'PNG') or os.path.join(output_folder, f"slide_{i}.png")
                for i, img in enumerate(convert_from_path(pdf_path, dpi=200))]

    if slides_data is None:
        raise RuntimeError("LibreOffice not found and no slide data provided for direct rendering.")

    print("LibreOffice not found — rendering slides with Pillow.")
    return _render_slides_pillow(slides_data, theme_colors, output_folder)


def _render_slides_pillow(slides_data, theme_colors, output_folder):
    from PIL import Image, ImageDraw, ImageFont

    if theme_colors is None:
        theme_colors = {"primary": "1F497D", "secondary": "4F81BD", "accent": "C0504D", "background": "FFFFFF", "text": "000000"}

    def _parse(c):
        c = c.lstrip("#")
        return tuple(int(c[i:i+2], 16) for i in (0, 2, 4))

    bg = _parse(theme_colors.get("background", "FFFFFF"))
    primary = _parse(theme_colors.get("primary", "1F497D"))
    secondary = _parse(theme_colors.get("secondary", "4F81BD"))
    accent = _parse(theme_colors.get("accent", "C0504D"))
    text_c = _parse(theme_colors.get("text", "000000"))

    W, H = 1920, 1080
    font_path = None
    for candidate in ["C:\\Windows\\Fonts\\segoeui.ttf", "C:\\Windows\\Fonts\\arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]:
        if os.path.exists(candidate):
            font_path = candidate
            break

    def _font(size, bold=False):
        if font_path:
            try:
                return ImageFont.truetype(font_path, size)
            except:
                pass
        return ImageFont.load_default()

    paths = []

    # Helper to wrap text
    def wrap_text(draw, text, font, max_width):
        words = text.split()
        lines = []
        line = ""
        for w in words:
            test = (line + " " + w).strip()
            bb = draw.textbbox((0, 0), test, font=font)
            if bb[2] - bb[0] <= max_width:
                line = test
            else:
                if line:
                    lines.append(line)
                line = w
        if line:
            lines.append(line)
        return lines

    for idx, slide in enumerate(slides_data):
        img = Image.new("RGB", (W, H), bg)
        draw = ImageDraw.Draw(img)

        # Header bar
        draw.rectangle([(0, 0), (W, 120)], fill=primary)

        # Slide title
        title_font = _font(44)
        title = slide.title if hasattr(slide, "title") else f"Slide {idx+1}"
        draw.text((60, 30), title, fill="white", font=title_font)

        # Content
        content = slide.content if hasattr(slide, "content") else ""
        content_font = _font(28)
        y = 170
        for line in wrap_text(draw, content, content_font, W - 120):
            draw.text((60, y), line, fill=text_c, font=content_font)
            y += 40
            if y > H - 100:
                break

        # Key points
        kp = slide.key_points if hasattr(slide, "key_points") else []
        y += 20
        bullet_font = _font(24)
        for pt in kp:
            draw.text((80, y), f"  {pt}", fill=secondary, font=bullet_font)
            y += 36
            if y > H - 80:
                break

        # Footer bar
        draw.rectangle([(0, H - 40), (W, H)], fill=primary)

        path = os.path.join(output_folder, f"slide_{idx}.png")
        img.save(path)
        paths.append(path)

    # Thank You slide
    img = Image.new("RGB", (W, H), bg)
    draw = ImageDraw.Draw(img)
    draw.rectangle([(0, 0), (W, 120)], fill=primary)
    ty_font = _font(60)
    ty_bb = draw.textbbox((0, 0), "Thank You!", font=ty_font)
    draw.text(((W - (ty_bb[2] - ty_bb[0])) // 2, H // 2 - 60), "Thank You!", fill=primary, font=ty_font)
    sub_font = _font(36)
    sub_bb = draw.textbbox((0, 0), "Please Like, Share And Subscribe?", font=sub_font)
    draw.text(((W - (sub_bb[2] - sub_bb[0])) // 2, H // 2 + 40), "Please Like, Share And Subscribe?", fill=accent, font=sub_font)
    draw.rectangle([(0, H - 40), (W, H)], fill=primary)
    path = os.path.join(output_folder, f"slide_{len(slides_data)}.png")
    img.save(path)
    paths.append(path)

    return paths
